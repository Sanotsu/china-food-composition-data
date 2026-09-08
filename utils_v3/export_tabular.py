#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""增强版数据表格导出：json_data_v3_*_fixed_en → 单个CSV + 单个XLSX

用途：为大模型 / 机器学习提供表格形态的完整数据。

输出（写入基准数据目录内）：
  food_composition_full.csv   全部 1677 条单文件（UTF-8 BOM，Excel 可直接打开不乱码）
  food_composition_full.xlsx  Sheet「全部数据」= 全部记录；
                              其后每个 merged_*.json 一个 Sheet（Sheet名=类别名）

原则：
  - 所有单元格一律按文本写入，原书符号与数值字符串逐字符保真
    （"2.50" 不变成 2.5、"Tr"/"—"/"(0)"/"091101x" 原样保留）
  - 首列为 category（来自文件名的类别名），其余列顺序与 JSON 字段顺序一致
  - 行序：按类别文件顺序（类别内 = foodCode 升序）

用法：
  python utils_v3/export_tabular.py                    # 默认基准目录
  python utils_v3/export_tabular.py --dir <数据目录>    # 指定其他基准目录
"""
import argparse
import csv
import json
import re
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_DIR = ROOT / "json_data_v3_20260825_qwen38max_kimi_k3_fixed_en"
CSV_NAME = "food_composition_full.csv"
XLSX_NAME = "food_composition_full.xlsx"
FULL_SHEET = "全部数据"


def load_records(base_dir: Path):
    """按类别文件顺序加载全部记录，category 取文件名（merged_ 之后部分）"""
    files = sorted(base_dir.glob("merged_*.json"))
    if not files:
        raise SystemExit(f"基准目录无 merged_*.json: {base_dir}")
    rows, columns = [], None
    for f in files:
        category = f.stem[len("merged_"):]
        records = json.loads(f.read_text(encoding="utf-8"))
        for rec in records:
            if columns is None:
                columns = list(rec.keys())
            elif list(rec.keys()) != columns:
                raise SystemExit(f"字段顺序不一致: {f.name} {rec.get('foodCode')}")
            rows.append({"category": category, **rec})
    return rows, ["category"] + columns


def safe_sheet_name(name: str) -> str:
    """Excel Sheet 名限制：禁用 :\\/?*[]，最长31字符"""
    return re.sub(r"[:\\/?*\[\]]", "_", name)[:31] or "sheet"


def export_csv(rows, columns, out_path: Path):
    with open(out_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def export_xlsx(rows, columns, out_path: Path):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    def write_sheet(ws, data):
        ws.append(columns)
        for rec in data:
            ws.append([str(rec.get(c, "")) for c in columns])
        # 文本格式 + 列宽（按内容自适应，封顶）
        widths = [len(c) * 2 for c in columns]
        for r_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            for c_idx, cell in enumerate(row, start=1):
                cell.number_format = "@"
                l = len(str(cell.value or ""))
                if l > widths[c_idx - 1]:
                    widths[c_idx - 1] = l
        for c_idx, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(c_idx)].width = min(w + 2, 60)
        ws.freeze_panes = "A2"

    write_sheet(wb.active, rows)
    wb.active.title = FULL_SHEET

    by_cat: dict = {}
    for rec in rows:
        by_cat.setdefault(rec["category"], []).append(rec)
    for category in sorted(by_cat):
        write_sheet(wb.create_sheet(safe_sheet_name(category)), by_cat[category])

    wb.save(out_path)
    return len(by_cat) + 1


def main():
    parser = argparse.ArgumentParser(description="增强版营养数据导出为 CSV/XLSX 表格（文本保真）")
    parser.add_argument("--dir", default=str(DEFAULT_DIR), help="基准数据目录（_fixed_en）")
    args = parser.parse_args()
    base_dir = Path(args.dir)

    rows, columns = load_records(base_dir)
    csv_path = base_dir / CSV_NAME
    xlsx_path = base_dir / XLSX_NAME

    export_csv(rows, columns, csv_path)
    sheets = export_xlsx(rows, columns, xlsx_path)

    en_cnt = sum(1 for r in rows if r.get("englishName", "") != "")
    print(f"基准: {base_dir}")
    print(f"记录: {len(rows)} 条（含英文名 {en_cnt} 条），列 {len(columns)} 列")
    print(f"已导出: {csv_path}")
    print(f"已导出: {xlsx_path}（{sheets} 个Sheet：{FULL_SHEET} + 每类别一个）")
    print(f"导出时间: {datetime.now().isoformat(timespec='seconds')}")


if __name__ == "__main__":
    main()
